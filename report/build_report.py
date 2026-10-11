"""Create an editable evidence-based report. Run with the bundled document Python."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'report';OUT.mkdir(exist_ok=True)
EVIDENCE=ROOT/'docs'/'evidence';SHOTS=EVIDENCE/'screenshots'
verified=json.loads((EVIDENCE/'environment.json').read_text(encoding='utf-8'))
assert verified['failures']==verified['errors']==verified['skipped']==0
assert verified['tests']==73
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.85)
sec.header_distance=sec.footer_distance=Inches(.3)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(9)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(28);doc.styles['Title'].font.bold=True
title_border=doc.styles['Title'].element.get_or_add_pPr().find(qn('w:pBdr'))
if title_border is not None:title_border.getparent().remove(title_border)
doc.styles['Heading 1'].font.size=Pt(22);doc.styles['Heading 1'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 2'].font.size=Pt(13);doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(10)
header=sec.header.paragraphs[0];header.text='PYTHON APPLICATIONS  /  INDEPENDENT DEMONSTRATIONS';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('526274')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Python application report  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('526274')
doc.core_properties.title='Design and development of Python applications'
doc.core_properties.subject='GUI media downloader and LAN multiplayer mathematics game'
doc.core_properties.author='[Student name]';doc.core_properties.keywords='Python, SunDrop, MathDuel, independent demonstration'

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def table(headers,rows,widths):
    tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for cell,text,width in zip(tab.rows[0].cells,headers,widths):cell.text=text;cell.width=Inches(width)
    repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=tab.add_row().cells
        for cell,text,width in zip(cells,row,widths):cell.text=str(text);cell.width=Inches(width)
    for i,row in enumerate(tab.rows):
        cannot=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cannot)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            prop=cell._tc.get_or_add_tcPr();fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'243B55' if i==0 else ('F1F4F7' if i%2 else 'FFFFFF'));prop.append(fill)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                line=OxmlElement('w:'+edge);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'D9D9D9');borders.append(line)
            prop.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'95');margin.set(qn('w:type'),'dxa');margins.append(margin)
            prop.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(10);run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '17232F')
    p('')
    return tab

def image(path,caption,max_width=6.5,max_height=6.2):
    with Image.open(path) as im:w,height=im.size
    width=min(max_width,max_height*w/height)
    para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(6)
    para.add_run().add_picture(str(path),width=Inches(width))
    cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Original diagram asset generated from the same architecture as the SVG.
diagram=Image.new('RGB',(1400,690),'#f7f8fa');draw=ImageDraw.Draw(diagram)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24);small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
boxes=[(50,100,330,200,'Laptop browser','Local host or LAN'),(50,260,330,360,'Phone browser','Intended deployment'),(500,155,850,310,'Flask and Socket.IO','Authoritative match engine'),(1050,180,1350,290,'SQLite','Completed match records'),(50,500,330,610,'SunDrop desktop','CustomTkinter interface'),(500,480,850,630,'Worker and yt-dlp','Artwork and safe downloads'),(1050,500,1350,610,'Local files and JSON','Library and thumbnail paths')]
for x,y,x2,y2,label,detail in boxes:
    draw.rounded_rectangle((x,y,x2,y2),radius=15,fill='white',outline='#62748a',width=3)
    draw.text(((x+x2)/2,(y+y2)/2-20),label,font=font,anchor='mm',fill='#243B55');draw.text(((x+x2)/2,(y+y2)/2+19),detail,font=small,anchor='mm',fill='#526274')
for pts in [((330,150),(425,150),(425,200),(500,200)),((330,310),(425,310),(425,260),(500,260)),((850,232),(1050,232)),((330,555),(500,555)),((850,555),(1050,555))]:draw.line(pts,fill='#62748a',width=4)
draw.text((50,35),'Independent applications with local storage',font=font,fill='#243B55')
diagram_path=OUT/'architecture.png';diagram.save(diagram_path)

# 1 Cover
p('ACADEMIC PROJECT REPORT').runs[0].font.size=Pt(10)
p('Design and development of Python applications','Title')
p('GUI media downloader and LAN multiplayer mathematics game','Subtitle')
p('SunDrop and MathDuel LAN').runs[0].font.size=Pt(17)
p('')
table(['Student information','Editable details'],[
('Name','[Student name]'),('Roll or enrolment number','[Roll number]'),('College and department','[College / department]'),('Course and academic year','[Course / year]'),('Faculty guide','[Faculty guide]'),('Industry mentor','[Mentor name if verified]')],[2.2,4.25])
p('Training details supplied in the project brief: Anishree Technologies LLP, Nagpur; Python; 15 June to 15 September 2026. Certificate date stated in the brief: 15 September 2026. These details require comparison with the original certificate before submission.')
p('Prepared on 11 October 2026. The applications are independent demonstrations developed for this report; their completion during the internship is not documented.')


# 2 Preliminary statements and contents
page('Acknowledgement declaration and abstract')
h('Acknowledgement and declaration')
p('I acknowledge the maintainers of the open-source libraries used in these demonstrations. This report describes the supplied source, tests and observed results. It does not certify attendance, company project work or completion during the internship. I will verify the training documents, understand the code and replace the student placeholders before submission.')
p('Student signature: [To be signed]     Date: [Submission date]')
h('Abstract')
p('SunDrop demonstrates authorized media downloading through a responsive Python desktop interface, with saved cover images, MP3 track details and a local audio library. MathDuel LAN demonstrates solo practice and two-player mental math with server-controlled questions, deadlines and scoring. The recorded suite contains 73 passing tests. Original audio/video fixtures and a two-browser-tab match were verified. The final SunDrop interface was captured from the running app. MathDuel screenshots record an earlier working version and are labelled. A physical phone networking test is still pending.')
h('Contents')
table(['Section','Page'],[
('Training overview and objectives',3),('Requirements and architecture',4),
('SunDrop implementation and interface',5),('MathDuel implementation and results',6),
('LAN deployment and mobile layout',7),('Testing and actual outcomes',8),
('Limitations learning outcomes and conclusion',9),('References and installation',10)],[5.8,.65])
p('This edition contains ten pages including the cover. Update contents and pagination after editing the DOCX.')

# 3 Training overview and objectives
page('Training overview and objectives')
p('The brief identifies Anishree Technologies LLP, Nagpur; Python training; 15 June to 15 September 2026; and certificate date 15 September 2026. The certificate and diary were not attached, so these particulars and the exact weekly syllabus remain unverified. The apps are independent demonstrations prepared on 11 October 2026, not documented company projects.')
table(['Week','Actual diary heading','Evidence'],[(str(i),'[Insert verified heading]','[Diary page / date]') for i in range(1,13)],[.65,3.7,2.1])
h('Objectives and educational connection')
p('Build a safe, responsive desktop media workflow and a fair local math game; separate interface code from rules and storage; and document reproducible tests. The code demonstrates types, functions, containers, exceptions, classes, threads, JSON, Flask, Socket.IO and SQLite. The actual diary must establish which topics were studied each week. BEGINNER_GUIDE.md, VIVA_PREP.md and DIARY_MAPPING.md provide the detailed learning material.')

# 4 Requirements and architecture
page('Requirements and architecture')
table(['Area','Implementation'],[
('Runtime','Python 3.11+; verified with 3.12.14 on Windows 11'),
('Desktop','CustomTkinter, yt-dlp, Pillow, Mutagen and pygame-ce'),
('Web and storage','Flask, Flask-SocketIO, HTML/CSS/JS, SQLite and JSON'),
('Optional media tool','FFmpeg for merging and selected MP3 conversion'),
('Permissions','Authorized media only; trusted LAN; no automatic firewall changes')],[1.3,5.15])
image(diagram_path,'Figure 1. Intended architecture. Physical phone connectivity has not been tested.',max_height=3.0)
p('SunDrop separates widgets in app.py, downloading in core.py, artwork in artwork.py and audio playback in player.py. A background task handles network requests. A queue passes its results to the main interface, where Tk after() safely updates the screen. JSON retains the latest 100 completed items and artwork paths.')
p('MathDuel separates Flask/Socket.IO transport from the locked game engine. Python owns answers, deadlines measured by a clock that cannot jump when the system time changes and scores. Active rooms stay in memory for up to 30 minutes; completed anonymous match records use SQLite queries that keep data values separate from SQL commands and 30-day retention. The browser displays public state and submits an answer with its round and match ID.')

# 5 SunDrop implementation and screenshot
page('SunDrop implementation and interface')
p('The desktop validates HTTP(S) URLs without embedded credentials and requires confirmation of download permission. yt-dlp retrieves title, creator, duration and actual source formats. Separate format and source-quality selectors avoid promised resolutions. A worker reports progress while the interface remains responsive; the download stops when its next safe cancellation check is reached. Filenames are sanitized for Windows. FFmpeg is detected before offering merge or MP3 conversion.')
p('Pillow normalizes thumbnails limited in file size and image dimensions into a reusable JPEG cache. Missing art gets an original title/type cover; a cover image saved next to the media file and persisted path keep the library association. Mutagen writes MP3 title, artist, album and source cover to a temporary copy, verifies tags and duration, then replaces the original. MP4 stays unchanged and uses a separate cover image. WAV/MP3/OGG playback supports pause, previous/next, progress and volume; moved files are marked.')
image(SHOTS/'sundrop-retro.png','Figure 2. Current SunDrop desktop: cyan and cream panels, retro menus and media preview. Captured on 11 October 2026.',max_height=3.5)
p('The current design uses pastel cyan and cream, charcoal outlines, monospace headings, hard shadows, a functional menu, library cards and a floating player. Hidden Tk tests cover responsiveness, library recovery, selectors, themes and playback navigation. The screenshot above was captured from the running final interface, using a separate empty library to protect personal download history.')

# 6 MathDuel implementation and results
page('MathDuel implementation and results')
p('The host creates a collision-checked six-digit room, limited to two players. A session token supports reconnecting the same tab. A duel has ten shared questions and up to 15 seconds per round, with two seconds of feedback. Addition, subtraction, multiplication, exact division, squares and integer roots produce whole-number answers at three difficulty levels.')
p('Python accepts one valid answer per player per round. Correct scoring is 10 + floor(seconds remaining / 3), clamped to a five-point bonus; incorrect or missing responses earn zero. Late, duplicate and old-match packets are rejected. Host rematch resets state; ties are allowed. Solo changes difficulty every five rounds using recent accuracy and speed rules, not machine learning.')
image(SHOTS/'mathduel-results.jpg','Figure 3. Actual completed two-tab match on one laptop, before final UI refinement. This is a local simulation.',max_width=4.1,max_height=4.2)
p('Ada scored 54 and Alan 53; both had 40% accuracy, with mean submitted times of 2.39 and 2.71 seconds. Several unattended rounds timed out. Host rematch and browser reload reconnection worked. The refined charcoal/lime UI retains original branding and mascot, adding mode navigation, round dots and copyable room codes. Its final visual walkthrough remains pending.')

# 7 Network deployment and mobile evidence
page('LAN deployment and mobile layout')
image(SHOTS/'mathduel-mobile-solo.jpg','Figure 4. Solo game in a 390 by 844 browser viewport. This is viewport simulation, not a physical phone test.',max_width=2.1,max_height=5.5)
p('The Windows laptop hosts Python on 0.0.0.0 with debug off. Both browsers must first open http://HOST_LAN_IP:PORT on a reachable local network, then enter a room code. The code selects a room; it does not discover the host. Check the active Wi-Fi IPv4 address and /health before troubleshooting joining.')
p('A phone may provide the hotspot while the laptop hosts, but hotspot or guest-network isolation can block traffic. Candidate addresses may include virtual adapters. Firewall Private-profile allowance is a user decision; no rules were changed. There is no public exposure or automatic phone hosting. Physical Android/iOS joining and experimental Termux hosting remain TODOs. The tested mobile viewport had 12 keypad controls and no horizontal overflow, before the final style refinement.')

# 8 Tests and evidence
page('Testing and actual outcomes')
p('Recorded run on 11 October 2026: 73 passed, zero failures, zero errors and zero skipped cases. Python 3.12.14 on Windows 11. JUnit XML, terminal output, environment versions, screenshot hashes and the browser record accompany the source.')
table(['Suite','Cases','Verified scope','Result'],[
('MathDuel engine and events','36','Generators, scoring, deadlines, duplicates, rooms, reconnect, rate limits and SQLite','Passed'),
('SunDrop core','23','URLs, safe names, history, cancellation and extraction errors','Passed'),
('Original WAV download','1','Real yt-dlp transfer; identical original/output bytes','Passed'),
('Artwork audio and MP4','10','Cache, fallback, MP3 details and embedded cover art, unchanged audio, conversion, MP4 cover image, opening dispatch and player','Passed'),
('Hidden Tk integration','3','Worker responsiveness, missing-file library, player navigation, format groups and themes','Passed')],[1.55,.5,3.5,.9])
h('Integration evidence')
p('Local HTTP fixtures served our own generated audio, video and artwork. MP3 title, artist, album and embedded cover art were read after saving; decoded audio hashes remained equal. A simulated tag-write failure preserved the original file. Original MP4 bytes stayed identical, with an associated separate cover image. Audio controls used a dummy device, so physical speaker output is not proven.')
p('Two real browser tabs shared questions and completed a match. Different answer arrival times produced different speed bonuses. Rematch and tab reload restored correct state. Engine tests use a controlled clock for deadlines; the browser demonstration used the real default duration. Latest JavaScript syntax and live routes were checked after refinement, but fresh layout capture remains pending.')
h('Evidence boundaries')
p('Passing cases prove their tested behaviors, not every source site or network. No physical phone, hotspot, Termux host, speaker-output check or actual OS-handler launch was verified. The current SunDrop screenshot and earlier MathDuel screenshots are labelled by version. No attendance, signatures or undocumented internship activities are asserted.')

# 9 Limitations and learning
page('Limitations learning outcomes and conclusion')
h('Limitations')
p('Media sites and terms can change. SunDrop does not guarantee all YouTube or Instagram links and does not bypass DRM, login or access controls. Cancellation can wait for a network call or processing step. Artwork cache pruning is not implemented; moving media without its cover leaves recorded paths missing. Supported internal playback is limited to WAV, MP3 and OGG.')
p('MathDuel uses one process and in-memory rooms; restarting loses active games. Local HTTP is unencrypted, room codes are not strong authentication and rate limits are basic. Disconnects reserve slots until expiry while clocks continue; there is no lobby kick button. Server scoring does not prevent outside calculators. The current SunDrop design has been captured. MathDuel’s final design and physical LAN operation need further checks.')
p('Student details are editable placeholders. Organization, internship dates, certificate wording and the 12 diary headings need their original documents. These demonstrations are not evidence of company work or completion during the internship.')
h('Learning outcomes and future scope')
p('The projects connect basic Python to visible applications: functions validate inputs, dictionaries hold rooms, classes organize state, exceptions support retryable errors, and files preserve records. Threads and queues explain responsiveness; Flask routes and Socket.IO show live communication. Parameterized queries and unique keys help maintain reliable statistics.')
p('Tests illustrate how to examine failures, boundary times and persistence with controlled fixtures. Original media avoids relying on copyrighted downloads or unstable public sites. The beginner guide contains runnable examples, and the viva guide contains 50 short answers; diary mapping must be added from verified material.')
p('Next steps are a documented physical-phone LAN test, a final MathDuel screenshot walkthrough, speaker and system-handler checks, optional QR join links, abandoned-slot recovery, accessibility review, cache retention and stronger deployment controls if the intended audience expands.')
h('Conclusion')
p('SunDrop and MathDuel implement their main independent demonstration workflows and pass the recorded 73-case suite. The ten-page report preserves architecture, the current SunDrop capture and labelled earlier MathDuel captures, results and clear limitations. It provides a practical basis for student explanation and further verification without claiming undocumented internship activity.')

# 10 References and concise installation appendix
page('References and installation')
p('Primary documentation consulted on 11 October 2026. The supplied brief is the sole source for unverified training particulars. Local evidence: docs/evidence/tests.xml, test-output.txt, environment.json, browser-check.json and screenshot-manifest.json.')
references=[
('Python Software Foundation','Python tutorial','https://docs.python.org/3/tutorial/'),
('Pallets','Flask documentation','https://flask.palletsprojects.com/en/stable/'),
('Flask SocketIO contributors','Getting started and deployment','https://flask-socketio.readthedocs.io/en/stable/'),
('yt dlp contributors','Official README','https://github.com/yt-dlp/yt-dlp'),
('CustomTkinter contributors','Documentation','https://customtkinter.tomschimansky.com/documentation/'),
('Python Software Foundation','sqlite3 reference','https://docs.python.org/3/library/sqlite3.html'),
('pytest contributors','Testing documentation','https://docs.pytest.org/en/stable/'),
('FFmpeg developers','Documentation','https://ffmpeg.org/documentation.html'),
('Mutagen contributors','ID3 guide','https://mutagen.readthedocs.io/en/latest/user/id3.html'),
('pygame community edition','Music reference','https://pyga.me/docs/ref/music.html'),
('Socket.IO contributors','Client API','https://socket.io/docs/v4/client-api/')]
for i,(author,title,url) in enumerate(references,1):
    para=p(f'[{i}] {author}. {title}.\n{url}')
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.line_spacing=1.0
    for run in para.runs:run.font.size=Pt(9.5)
h('PowerShell quick start')
p('From the project folder, use Python 3.11+ and an isolated environment. Run the apps in separate terminals. Complete startup, firewall troubleshooting and optional FFmpeg notes are in README.md.')
commands=[
'py -3.12 -m venv .venv',
'.\\.venv\\Scripts\\python.exe -m pip install -r mathduel-lan\\requirements.txt -r sundrop\\requirements.txt',
'.\\.venv\\Scripts\\python.exe mathduel-lan\\app.py',
'.\\.venv\\Scripts\\python.exe sundrop\\app.py',
'.\\.venv\\Scripts\\python.exe -m pytest mathduel-lan\\tests sundrop\\tests -q']
for command in commands:
    para=p(command);para.paragraph_format.space_after=Pt(4)
    for run in para.runs:run.font.name='Consolas';run.font.size=Pt(8.5)
p('Open http://127.0.0.1:5000 on the host; use its LAN IPv4 URL on a phone. Choose --port 5001 if needed. README.md includes the original-tone demo and GitHub commands. GitHub publication and live hosting are being prepared; this report does not claim a deployment that has not been verified. Review student placeholders and evidence before publication; editing this report requires a fresh ten-page pagination check.')

path=OUT/'Python_Applications_Report.docx';doc.save(path)
print(path)
